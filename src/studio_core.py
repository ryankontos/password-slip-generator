"""Workbook import, rule evaluation, and PDF rendering for Password Slip Studio."""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Iterable

from openpyxl import load_workbook
from reportlab.lib.colors import Color, HexColor
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.pdfbase.pdfmetrics import getAscentDescent, stringWidth
from reportlab.pdfgen import canvas


MM = 72 / 25.4
MAX_IMPORT_ROWS = 20_000
MAX_IMPORT_COLUMNS = 200
PAPER_SIZES = {"a4": A4, "letter": LETTER}


class StudioError(ValueError):
    """A user-facing validation failure."""


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    return str(value)


def unique_headers(values: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(values, start=1):
        base = clean_cell(raw).strip() or f"Column {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return headers


def parse_workbook(filename: str, content: bytes) -> dict[str, Any]:
    """Return every visible worksheet as JSON-friendly headers and rows."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        return {"filename": filename, "sheets": [_parse_csv(content)]}
    if suffix not in {"xlsx", "xlsm"}:
        raise StudioError("Choose a .xlsx, .xlsm, or .csv file.")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise StudioError("That workbook could not be read.") from exc
    sheets = []
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            rows: list[list[str]] = []
            row_numbers: list[int] = []
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, ())
            headers = unique_headers(first[:MAX_IMPORT_COLUMNS])
            for row_number, row in enumerate(iterator, start=2):
                dimension = sheet.row_dimensions.get(row_number)
                if dimension and dimension.hidden:
                    continue
                values = [clean_cell(value) for value in row[: len(headers)]]
                if any(value.strip() for value in values):
                    rows.append(values)
                    row_numbers.append(row_number)
                if len(rows) >= MAX_IMPORT_ROWS:
                    break
            sheets.append({
                "name": sheet.title,
                "headers": headers,
                "rows": rows,
                "rowNumbers": row_numbers,
                "truncated": len(rows) >= MAX_IMPORT_ROWS,
            })
    finally:
        workbook.close()
    if not sheets:
        raise StudioError("The workbook has no visible sheets.")
    return {"filename": filename, "sheets": sheets}


def _parse_csv(content: bytes) -> dict[str, Any]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise StudioError("The CSV text encoding could not be read.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    first = next(reader, [])
    headers = unique_headers(first[:MAX_IMPORT_COLUMNS])
    rows = []
    row_numbers = []
    for row_number, row in enumerate(reader, start=2):
        values = [clean_cell(value) for value in row[: len(headers)]]
        values.extend([""] * (len(headers) - len(values)))
        if any(value.strip() for value in values):
            rows.append(values)
            row_numbers.append(row_number)
        if len(rows) >= MAX_IMPORT_ROWS:
            break
    return {"name": "CSV", "headers": headers, "rows": rows, "rowNumbers": row_numbers, "truncated": len(rows) >= MAX_IMPORT_ROWS}


def condition_matches(condition: dict[str, Any], values: dict[str, Any]) -> bool:
    actual = clean_cell(values.get(str(condition.get("field", "")), ""))
    expected = clean_cell(condition.get("value", ""))
    operator = str(condition.get("operator", "not_empty"))
    left = actual.casefold()
    right = expected.casefold()
    if operator == "not_empty":
        return bool(actual.strip())
    if operator == "empty":
        return not actual.strip()
    if operator == "equals":
        return left == right
    if operator == "not_equals":
        return left != right
    if operator == "contains":
        return right in left
    if operator == "not_contains":
        return right not in left
    if operator == "starts_with":
        return left.startswith(right)
    if operator == "ends_with":
        return left.endswith(right)
    if operator == "greater_than":
        try:
            return float(actual) > float(expected)
        except ValueError:
            return False
    if operator == "less_than":
        try:
            return float(actual) < float(expected)
        except ValueError:
            return False
    if operator == "matches":
        try:
            return bool(re.search(expected, actual, flags=re.IGNORECASE))
        except re.error:
            return False
    return False


def rule_matches(rule: dict[str, Any], values: dict[str, Any]) -> bool:
    conditions = rule.get("conditions")
    if not isinstance(conditions, list) or not conditions:
        return False
    results = [condition_matches(item, values) for item in conditions if isinstance(item, dict)]
    if not results:
        return False
    matched = any(results) if rule.get("match") == "any" else all(results)
    return not matched if rule.get("negate") else matched


def included_rows(state: dict[str, Any]) -> list[dict[str, Any]]:
    """Return printable rows after explicit row state and hide-slip rules."""
    hide_rules = [
        rule for rule in state.get("rules", [])
        if isinstance(rule, dict) and rule.get("enabled", True) and rule.get("action") == "hide_slip"
    ]
    result = []
    for row in state.get("rows", []):
        if not isinstance(row, dict) or row.get("hidden") or row.get("disabled"):
            continue
        values = row.get("values", {})
        if any(rule_matches(rule, values) for rule in hide_rules):
            continue
        result.append(row)
    return result


def visible_columns(state: dict[str, Any], row: dict[str, Any]) -> list[dict[str, Any]]:
    values = row.get("values", {})
    overrides = row.get("overrides", {})
    active_rules = [rule for rule in state.get("rules", []) if rule.get("enabled", True)]
    visible = []
    for column in state.get("columns", []):
        column_id = str(column.get("id", ""))
        mode = column.get("visibility", "always")
        # Field visibility is authoritative. A sheet-level layout preference
        # must never turn an empty "Only with a value" field back on.
        value_has_content = any(character.isalnum() for character in clean_cell(values.get(column_id, "")))
        shown = mode == "always" or (mode == "nonempty" and value_has_content)
        matching_actions = {
            rule.get("action") for rule in active_rules
            if rule.get("target") == column_id and rule_matches(rule, values)
        }
        # A matching hide always wins over a matching show. This makes rule
        # order irrelevant and prevents a later rule from revealing a field.
        if "hide_field" in matching_actions:
            shown = False
        elif "show_field" in matching_actions:
            shown = True
        if overrides.get(column_id) is True:
            shown = True
        elif overrides.get(column_id) is False:
            shown = False
        if shown:
            visible.append(column)
    return visible


@dataclass
class PdfLayout:
    mode: str
    page_size: tuple[float, float]
    margin: float
    gap: float
    slip_height: float
    accent: HexColor
    ink: HexColor
    muted: HexColor
    border: HexColor
    paper_color: HexColor
    label_size: float
    value_size: float
    font: str
    label_font: str
    label_case: str
    value_align: str
    label_width: float
    padding: float
    show_border: bool
    cut_marks: bool
    footer: bool
    field_lines: bool
    stacked_columns: int


def _number(value: Any, default: float, low: float, high: float) -> float:
    try:
        return min(high, max(low, float(value)))
    except (TypeError, ValueError):
        return default


def _hex(value: Any, default: str) -> HexColor:
    try:
        return HexColor(str(value))
    except (TypeError, ValueError):
        return HexColor(default)


def pdf_layout(state: dict[str, Any]) -> PdfLayout:
    source = state.get("layout", {})
    page = PAPER_SIZES.get(str(source.get("paper", "a4")), A4)
    if source.get("orientation") == "landscape":
        page = (page[1], page[0])
    return PdfLayout(
        mode="stacked" if source.get("mode") == "stacked" else "horizontal",
        page_size=page,
        margin=_number(source.get("margin", 10), 10, 2, 40) * MM,
        gap=_number(source.get("gap", 0), 0, 0, 20) * MM,
        slip_height=_number(source.get("slipHeight", 36), 36, 20, 120) * MM,
        accent=_hex(source.get("accent"), "#00539B"),
        ink=_hex(source.get("ink"), "#151719"),
        muted=_hex(source.get("muted"), "#687078"),
        border=_hex(source.get("borderColor"), "#C9CED4"),
        paper_color=_hex(source.get("paperColor"), "#FFFFFF"),
        label_size=_number(source.get("labelSize", 10), 10, 4, 18),
        value_size=_number(source.get("valueSize", 14), 14, 5, 26),
        font=str(source.get("font", "Helvetica")) if source.get("font") in {"Helvetica", "Times-Roman", "Courier"} else "Helvetica",
        label_font=str(source.get("labelFont", source.get("font", "Helvetica"))) if source.get("labelFont", source.get("font", "Helvetica")) in {"Helvetica", "Times-Roman", "Courier"} else "Helvetica",
        label_case=str(source.get("labelCase", "original")),
        value_align=str(source.get("valueAlign", "left")),
        label_width=_number(source.get("labelWidth", 34), 34, 18, 60) / 100,
        padding=_number(source.get("padding", 2), 2, 0, 10) * MM,
        show_border=bool(source.get("showBorder", False)),
        cut_marks=bool(source.get("cutMarks", True)),
        footer=bool(source.get("footer", True)),
        field_lines=bool(source.get("fieldLines", True)),
        stacked_columns=2 if source.get("stackedColumns") == 2 else 1,
    )


def row_pdf_layout(state: dict[str, Any], row: dict[str, Any], base: PdfLayout) -> PdfLayout:
    """Compatibility shim: all slips now use the single sheet layout."""
    return base


def _footer_baseline(layout: PdfLayout) -> float:
    """Place footer text inside, and relative to, the bottom page margin."""
    return max(2.5 * MM, layout.margin * 0.55)


def render_pdf(state: dict[str, Any]) -> bytes:
    rows = included_rows(state)
    if not rows:
        if state.get("rows"):
            raise StudioError("There are no printable rows. Check hidden row settings and hide-slip rules.")
        raise StudioError("There are no data rows to export. Import a sheet or add a row first.")
    columns = state.get("columns", [])
    if not columns:
        raise StudioError("Add at least one field before exporting.")
    layout = pdf_layout(state)
    width, height = layout.page_size
    footer_height = 7 * MM if layout.footer else 0
    usable_width = width - 2 * layout.margin
    slip_width = usable_width
    usable_height = height - 2 * layout.margin - footer_height
    down = int((usable_height + layout.gap) // (layout.slip_height + layout.gap))
    if down < 1:
        raise StudioError("The slip height and page margins do not fit on the selected paper.")
    per_page = down
    pages = math.ceil(len(rows) / per_page)
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, pagesize=layout.page_size, pageCompression=1)
    pdf.setTitle(str(state.get("name") or "Password Slip Studio"))
    exported_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")
    for page_index in range(pages):
        page_rows = rows[page_index * per_page : (page_index + 1) * per_page]
        for index, row in enumerate(page_rows):
            top = height - layout.margin - index * (layout.slip_height + layout.gap)
            _draw_slip(pdf, row, visible_columns(state, row), layout, layout.margin, top - layout.slip_height, slip_width, layout.slip_height)
        if layout.cut_marks:
            _draw_cut_marks(pdf, layout, width, height, slip_width, down)
        if layout.footer:
            pdf.setFillColor(layout.muted)
            pdf.setFont("Helvetica", 7)
            footer_y = _footer_baseline(layout)
            pdf.drawString(layout.margin, footer_y, f"Exported {exported_at}")
            pdf.drawRightString(width - layout.margin, footer_y, f"Page {page_index + 1} of {pages}")
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


def _draw_cut_marks(pdf: canvas.Canvas, layout: PdfLayout, width: float, height: float, slip_width: float, down: int) -> None:
    pdf.saveState()
    pdf.setStrokeColor(layout.border)
    pdf.setLineWidth(0.35)
    tick = 3 * MM
    for row_index in range(down + 1):
        y = height - layout.margin - row_index * (layout.slip_height + layout.gap)
        pdf.line(0, y, tick, y)
        pdf.line(width - tick, y, width, y)
    pdf.restoreState()


def _font_for(column: dict[str, Any], layout: PdfLayout) -> str:
    style = column.get("style")
    if style == "mono" or column.get("type") == "password":
        return "Courier-Bold"
    if style == "strong":
        return {"Helvetica": "Helvetica-Bold", "Times-Roman": "Times-Bold", "Courier": "Courier-Bold"}.get(layout.font, "Helvetica-Bold")
    return layout.font


def _label_font(layout: PdfLayout) -> str:
    return {"Helvetica": "Helvetica-Bold", "Times-Roman": "Times-Bold", "Courier": "Courier-Bold"}.get(layout.label_font, "Helvetica-Bold")


def _label(value: Any, layout: PdfLayout) -> str:
    text = clean_cell(value)
    if layout.label_case == "upper":
        return text.upper()
    if layout.label_case == "title":
        return text.title()
    return text


def _display_value(column: dict[str, Any], value: Any) -> str:
    text = clean_cell(value)
    transform = str(column.get("valueTransform", "as_entered"))
    if transform == "upper":
        return text.upper()
    if transform == "lower":
        return text.lower()
    if transform == "title":
        return text.title()
    if transform == "mask_last4" and len(text) > 4:
        return "•" * max(4, len(text) - 4) + text[-4:]
    return text


def _value_align(column: dict[str, Any], layout: PdfLayout) -> str:
    alignment = column.get("valueAlign")
    return alignment if alignment in {"left", "center", "right"} else layout.value_align


def _fit(text: str, font: str, size: float, width: float, minimum: float = 4) -> tuple[str, float]:
    if not text:
        return "", size
    measured = stringWidth(text, font, size)
    fitted = max(minimum, min(size, size * width / measured)) if measured else size
    if fitted > minimum or stringWidth(text, font, fitted) <= width:
        return text, fitted
    suffix = "…"
    while text and stringWidth(text + suffix, font, fitted) > width:
        text = text[:-1]
    return text + suffix, fitted


def _text(pdf: canvas.Canvas, text: Any, font: str, size: float, color: Color, x: float, baseline: float, width: float, align: str = "left") -> None:
    cleaned, fitted = _fit(clean_cell(text), font, size, max(1, width))
    pdf.setFillColor(color)
    pdf.setFont(font, fitted)
    if align == "center":
        pdf.drawCentredString(x + width / 2, baseline, cleaned)
    elif align == "right":
        pdf.drawRightString(x + width, baseline, cleaned)
    else:
        pdf.drawString(x, baseline, cleaned)


def _centered_text(pdf: canvas.Canvas, text: Any, font: str, size: float, color: Color, x: float, bottom: float, width: float, height: float, align: str = "left") -> None:
    """Draw fitted text optically centred within a field row."""
    cleaned, fitted = _fit(clean_cell(text), font, size, max(1, width))
    ascent, descent = getAscentDescent(font, fitted)
    baseline = bottom + (height - (ascent - descent)) / 2 - descent
    pdf.setFillColor(color)
    pdf.setFont(font, fitted)
    if align == "center":
        pdf.drawCentredString(x + width / 2, baseline, cleaned)
    elif align == "right":
        pdf.drawRightString(x + width, baseline, cleaned)
    else:
        pdf.drawString(x, baseline, cleaned)


def _draw_slip(pdf: canvas.Canvas, row: dict[str, Any], columns: list[dict[str, Any]], layout: PdfLayout, x: float, y: float, width: float, height: float) -> None:
    pdf.saveState()
    pdf.setFillColor(layout.paper_color)
    pdf.setStrokeColor(layout.border)
    pdf.setLineWidth(0.6)
    pdf.rect(x, y, width, height, fill=1, stroke=int(layout.show_border))
    values = row.get("values", {})
    if not columns:
        pdf.setFillColor(layout.muted)
        pdf.setFont("Helvetica-Oblique", layout.value_size)
        pdf.drawCentredString(x + width / 2, y + height / 2, "No visible fields")
        pdf.restoreState()
        return
    if layout.mode == "stacked":
        _draw_stacked(pdf, values, columns, layout, x, y, width, height)
    else:
        _draw_horizontal(pdf, values, columns, layout, x, y, width, height)
    pdf.restoreState()


def _draw_horizontal(pdf: canvas.Canvas, values: dict[str, Any], columns: list[dict[str, Any]], layout: PdfLayout, x: float, y: float, width: float, height: float) -> None:
    label_height = min(height * 0.5, max(10 * MM, layout.label_size * 2.1))
    cell_width = width / len(columns)
    cursor = x
    for index, column in enumerate(columns):
        pdf.setFillColor(layout.accent)
        pdf.rect(cursor, y + height - label_height, cell_width, label_height, fill=1, stroke=0)
        padding = max(layout.padding, 1 * MM)
        _text(pdf, _label(column.get("label", ""), layout), _label_font(layout), layout.label_size, HexColor("#FFFFFF"), cursor + padding, y + height - label_height + (label_height - layout.label_size) / 2, cell_width - padding * 2, "center")
        _text(pdf, _display_value(column, values.get(column.get("id"), "")), _font_for(column, layout), layout.value_size, layout.ink, cursor + padding, y + (height - label_height - layout.value_size) / 2, cell_width - padding * 2, _value_align(column, layout))
        if index and layout.field_lines:
            pdf.setStrokeColor(layout.border)
            pdf.line(cursor, y, cursor, y + height)
        cursor += cell_width


def _draw_stacked(pdf: canvas.Canvas, values: dict[str, Any], columns: list[dict[str, Any]], layout: PdfLayout, x: float, y: float, width: float, height: float) -> None:
    blocks = layout.stacked_columns
    rows_per_block = math.ceil(len(columns) / blocks)
    block_width = width / blocks
    row_height = height / rows_per_block
    label_width = min(block_width * layout.label_width, 55 * MM)
    for index, column in enumerate(columns):
        block = index // rows_per_block
        row_index = index % rows_per_block
        left = x + block * block_width
        bottom = y + height - (row_index + 1) * row_height
        pdf.setFillColor(layout.accent)
        pdf.rect(left, bottom, label_width, row_height, fill=1, stroke=0)
        padding = max(layout.padding, 1 * MM)
        _centered_text(pdf, _label(column.get("label", ""), layout), _label_font(layout), layout.label_size, HexColor("#FFFFFF"), left + padding, bottom, label_width - padding * 2, row_height)
        _centered_text(pdf, _display_value(column, values.get(column.get("id"), "")), _font_for(column, layout), layout.value_size, layout.ink, left + label_width + padding, bottom, block_width - label_width - padding * 2, row_height, _value_align(column, layout))
    if layout.field_lines:
        # Draw dividers last so they remain continuous over the filled label
        # panel. The label-side segment is derived from the accent itself, so
        # it stays visible even when the configured rule colour matches it.
        label_divider = Color(
            layout.accent.red + (1 - layout.accent.red) * 0.58,
            layout.accent.green + (1 - layout.accent.green) * 0.58,
            layout.accent.blue + (1 - layout.accent.blue) * 0.58,
        )
        pdf.setLineWidth(0.6)
        for block in range(blocks):
            left = x + block * block_width
            count = min(rows_per_block, len(columns) - block * rows_per_block)
            for row_index in range(1, count):
                divider_y = y + height - row_index * row_height
                pdf.setStrokeColor(label_divider)
                pdf.line(left, divider_y, left + label_width, divider_y)
                pdf.setStrokeColor(layout.border)
                pdf.line(left + label_width, divider_y, left + block_width, divider_y)
            if block:
                pdf.setStrokeColor(layout.border)
                pdf.line(left, y, left, y + height)
