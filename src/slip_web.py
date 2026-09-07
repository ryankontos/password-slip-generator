#!/usr/bin/env python3
"""Local web workspace for editing and exporting password slips."""

from __future__ import annotations

import argparse
import base64
import binascii
from dataclasses import asdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
import json
import mimetypes
import os
from pathlib import Path
import re
import socket
import threading
import time
from typing import Any, Dict, Optional
from urllib.parse import parse_qs, urlparse
import webbrowser

from openpyxl import load_workbook

from password_slips import (
    APP_NAME,
    ROOT_DIR,
    Settings,
    column_letter,
    email_subject,
    layout_field_names,
    make_pdf,
    newest_workbook,
    open_default_application,
    open_email_draft,
    output_path,
    read_saved_settings,
    serialized_row_filter_sets,
    unique_labels,
)
from slip_preview import preview_page


WEB_ROOT = (ROOT_DIR / "web").resolve()
WEB_SETTINGS = ROOT_DIR / "settings" / "web_workspace.json"
WEB_UPLOAD = ROOT_DIR / "settings" / "web_workbook.xlsx"
TMP_DIR = ROOT_DIR / "tmp" / "web"
MAX_BODY = 40 * 1024 * 1024
MAX_UPLOAD = 25 * 1024 * 1024
LOOPBACK_HOSTS = {"127.0.0.1", "localhost", "::1"}
SAFE_FONTS = {
    "Helvetica", "Helvetica-Bold", "Helvetica-Oblique",
    "Times-Roman", "Times-Bold", "Times-Italic",
    "Courier", "Courier-Bold", "Courier-Oblique",
}


class WebError(ValueError):
    def __init__(self, message: str, status: int = 400) -> None:
        super().__init__(message)
        self.status = status


def _json_read(path: Path) -> Dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _json_write(path: Path, value: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    return str(value)


def _safe_filename(value: str, fallback: str) -> str:
    name = Path(value).name.strip()
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "", name).strip(" .")
    return cleaned[:120] or fallback


def _header_key(value: Any) -> str:
    return " ".join(str(value).casefold().split())


def _design_json(settings: Settings) -> Dict[str, Any]:
    return {name: getattr(settings, name) for name in layout_field_names()}


def _column_json(headers: list[str], settings: Settings) -> list[Dict[str, Any]]:
    name_order = {
        _header_key(name): index for index, name in enumerate(settings.default_column_names)
    }
    password_names = {_header_key(name) for name in settings.password_column_names}
    truncate_names = {_header_key(name) for name in settings.truncate_column_names}
    has_name_matches = any(_header_key(header) in name_order for header in headers)
    remembered = settings.column_numbers or list(range(1, len(headers) + 1))
    output = []
    for index, header in enumerate(headers):
        number = index + 1
        key = _header_key(header)
        output.append({
            "id": "column-%d" % number,
            "sourceIndex": index,
            "letter": column_letter(number),
            "sourceName": header,
            "label": header,
            "included": key in name_order if has_name_matches else (
                number in remembered if not name_order else True
            ),
            "password": key in password_names if password_names else number in settings.password_column_numbers,
            "truncate": key in truncate_names if truncate_names else number in settings.truncate_column_numbers,
            "defaultOrder": name_order.get(key, len(name_order) + index),
        })
    if has_name_matches:
        output.sort(key=lambda item: item["defaultOrder"])
    else:
        output.sort(key=lambda item: remembered.index(item["sourceIndex"] + 1)
                    if item["sourceIndex"] + 1 in remembered else len(remembered) + item["sourceIndex"])
    return output


class Application:
    def __init__(self) -> None:
        self.settings = read_saved_settings()
        self.workspace = _json_read(WEB_SETTINGS)
        self.preview_pdf: Optional[Path] = None
        self.export_pdf: Optional[Path] = None
        self.workbook_name = str(self.workspace.get("sourceName", ""))
        self.lock = threading.RLock()
        TMP_DIR.mkdir(parents=True, exist_ok=True)

    def state_json(self) -> Dict[str, Any]:
        latest = newest_workbook(self.settings.input_folder, self.settings.workbook_extensions)
        return {
            "workspace": self.workspace,
            "defaults": {
                "design": _design_json(self.settings),
                "outputFolder": self.settings.output_folder,
                "emailAddress": self.settings.email_address,
                "savedRuleSets": serialized_row_filter_sets(self.settings.saved_row_filter_sets),
                "fields": [
                    {
                        "name": name,
                        "password": _header_key(name) in {
                            _header_key(value) for value in self.settings.password_column_names
                        },
                        "truncate": _header_key(name) in {
                            _header_key(value) for value in self.settings.truncate_column_names
                        },
                    }
                    for name in self.settings.default_column_names
                ],
            },
            "latestWorkbook": str(latest) if latest else "",
            "fonts": sorted(SAFE_FONTS),
        }

    def save_workspace(self, value: Any) -> Dict[str, Any]:
        if not isinstance(value, dict):
            raise WebError("The workspace data was invalid.")
        encoded = json.dumps(value, ensure_ascii=False).encode("utf-8")
        if len(encoded) > MAX_BODY:
            raise WebError("This workspace is too large to save.", 413)
        self.workspace = value
        self.workbook_name = str(value.get("sourceName", ""))
        _json_write(WEB_SETTINGS, value)
        return {"saved": True, "savedAt": int(time.time())}

    def import_upload(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        name = _safe_filename(str(payload.get("name", "")), "workbook.xlsx")
        try:
            raw = base64.b64decode(str(payload.get("data", "")), validate=True)
        except (ValueError, binascii.Error) as exc:
            raise WebError("The dropped file could not be read.") from exc
        if not raw or len(raw) > MAX_UPLOAD:
            raise WebError("Choose a file smaller than 25 MB.", 413)
        if name.lower().endswith(".pdf"):
            path = TMP_DIR / "dropped-preview.pdf"
            path.write_bytes(raw)
            self.preview_pdf = path
            return {
                "kind": "pdf",
                "name": name,
                "url": "/api/file?kind=preview&t=%d" % time.time_ns(),
            }
        if not name.lower().endswith((".xlsx", ".xlsm")):
            raise WebError("Drop an Excel workbook or PDF.")
        WEB_UPLOAD.parent.mkdir(parents=True, exist_ok=True)
        WEB_UPLOAD.write_bytes(raw)
        self.workbook_name = name
        return self.read_workbook(
            name,
            str(payload.get("sheet", "")),
            bool(payload.get("excludeHidden", True)),
        )

    def import_latest(self) -> Dict[str, Any]:
        path = newest_workbook(self.settings.input_folder, self.settings.workbook_extensions)
        if not path:
            raise WebError("No Excel workbook was found in Downloads.", 404)
        raw = path.read_bytes()
        if len(raw) > MAX_UPLOAD:
            raise WebError("The newest workbook is larger than 25 MB.", 413)
        WEB_UPLOAD.parent.mkdir(parents=True, exist_ok=True)
        WEB_UPLOAD.write_bytes(raw)
        self.workbook_name = path.name
        return self.read_workbook(path.name, "", True)

    def switch_sheet(self, sheet: str, exclude_hidden: bool = True) -> Dict[str, Any]:
        if not WEB_UPLOAD.is_file():
            raise WebError("Drop the workbook again to switch sheets.", 404)
        return self.read_workbook(self.workbook_name or WEB_UPLOAD.name, sheet, exclude_hidden)

    def read_workbook(self, name: str, requested_sheet: str,
                      exclude_hidden: bool = True) -> Dict[str, Any]:
        try:
            workbook = load_workbook(WEB_UPLOAD, data_only=True)
        except Exception as exc:
            raise WebError("That workbook could not be opened.") from exc
        try:
            sheets = workbook.sheetnames
            if not sheets:
                raise WebError("That workbook has no worksheets.")
            sheet_name = requested_sheet if requested_sheet in sheets else sheets[0]
            sheet = workbook[sheet_name]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = unique_labels(header_row)
            rows = []
            hidden_count = 0
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                dimension = sheet.row_dimensions.get(row_number)
                hidden = bool(dimension and dimension.hidden)
                if hidden:
                    hidden_count += 1
                if hidden and exclude_hidden:
                    continue
                values = [_text(value) for value in row[:len(headers)]]
                values += [""] * (len(headers) - len(values))
                if any(value.strip() for value in values):
                    rows.append({
                        "id": "sheet-%d" % row_number,
                        "rowNumber": row_number,
                        "values": values,
                        "included": True,
                        "added": False,
                        "hidden": hidden,
                    })
            return {
                "kind": "workbook",
                "sourceName": name,
                "sheets": sheets,
                "sheet": sheet_name,
                "headers": headers,
                "columns": _column_json(headers, self.settings),
                "rows": rows,
                "hiddenRows": hidden_count,
                "excludeHidden": exclude_hidden,
            }
        finally:
            workbook.close()

    def settings_for_job(self, workspace: Dict[str, Any]) -> tuple[Settings, list[list[str]]]:
        columns = workspace.get("columns", [])
        rows = workspace.get("rowsToPrint", [])
        if not isinstance(columns, list) or not isinstance(rows, list):
            raise WebError("The slip data was invalid.")
        selected = [item for item in columns if isinstance(item, dict) and item.get("included")]
        if not selected:
            raise WebError("Choose at least one field.")

        settings = Settings()
        design = workspace.get("design", {})
        if isinstance(design, dict):
            for name in layout_field_names():
                if name in design:
                    from password_slips import clean_layout_value
                    setattr(settings, name, clean_layout_value(settings, name, design[name]))
        settings.columns = [str(item.get("label") or item.get("sourceName") or "Field") for item in selected]
        settings.column_numbers = list(range(1, len(selected) + 1))
        settings.password_column_numbers = [
            index for index, item in enumerate(selected, start=1) if item.get("password")
        ]
        settings.truncate_column_numbers = [
            index for index, item in enumerate(selected, start=1) if item.get("truncate")
        ]
        settings.sheet = str(workspace.get("sheet", "Workbook"))
        settings.workbook = str(workspace.get("sourceName", "password slips.xlsx"))
        settings.output_folder = str(workspace.get("outputFolder") or self.settings.output_folder)
        settings.email_address = str(workspace.get("emailAddress", "")).strip()
        settings.blank_slips = max(0, int(workspace.get("blankSlips", 0) or 0))
        extra = workspace.get("extraSummaryColumns", [])
        settings.extra_summary_columns = [str(value).strip() for value in extra if str(value).strip()] \
            if isinstance(extra, list) else []

        data = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            values = row.get("values", [])
            if not isinstance(values, list):
                values = []
            data.append([
                _text(values[int(item.get("sourceIndex", 0))])
                if 0 <= int(item.get("sourceIndex", 0)) < len(values) else ""
                for item in selected
            ])
        if not data and settings.blank_slips == 0:
            raise WebError("Include at least one row or add blank slips.")
        return settings, data

    def render(self, workspace: Dict[str, Any], export: bool) -> Dict[str, Any]:
        settings, rows = self.settings_for_job(workspace)
        with self.lock:
            if export:
                output = output_path(settings)
            else:
                output = TMP_DIR / "password-slips-preview.pdf"
            count, pages = make_pdf(settings, output, rows)
        if export:
            self.export_pdf = output
            kind = "export"
        else:
            self.preview_pdf = output
            kind = "preview"
        return {
            "kind": kind,
            "url": "/api/file?kind=%s&t=%d" % (kind, time.time_ns()),
            "downloadUrl": "/api/file?kind=%s&download=1&t=%d" % (kind, time.time_ns()),
            "path": str(output),
            "name": output.name,
            "slips": count,
            "pages": pages,
        }

    def preview(self, workspace: Dict[str, Any], page: int) -> Dict[str, Any]:
        settings, rows = self.settings_for_job(workspace)
        with self.lock:
            return preview_page(
                settings,
                rows,
                page,
                time.strftime("%Y-%m-%d %H:%M"),
            )

    def open_pdf(self, kind: str) -> Dict[str, Any]:
        path = self.export_pdf if kind == "export" else self.preview_pdf
        if not path or not path.is_file():
            raise WebError("Generate the PDF first.", 404)
        open_default_application(str(path))
        return {"opened": True, "path": str(path)}

    def email(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        address = str(payload.get("emailAddress", "")).strip()
        sheet = str(payload.get("sheet", "Workbook"))
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", address):
            raise WebError("Enter a valid email address.")
        if not open_email_draft(address, email_subject(Settings(sheet=sheet))):
            raise WebError("The default mail app could not open.", 500)
        return {"opened": True}


class Handler(BaseHTTPRequestHandler):
    server_version = "PasswordSlipGenerator/1.0"
    sys_version = ""

    @property
    def app(self) -> Application:
        return self.server.app  # type: ignore[attr-defined]

    def log_message(self, format: str, *args: Any) -> None:
        return

    def _same_local_server(self, value: str) -> bool:
        try:
            parsed = urlparse(value)
            return bool(parsed.hostname in LOOPBACK_HOSTS and parsed.port == self.server.server_port)
        except ValueError:
            return False

    def _allowed(self) -> bool:
        host = self.headers.get("Host", "")
        origin = self.headers.get("Origin", "")
        return self._same_local_server("http://" + host) and (not origin or self._same_local_server(origin))

    def _send(self, body: bytes, content_type: str, status: int = 200,
              headers: Optional[Dict[str, str]] = None) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "SAMEORIGIN")
        self.send_header("Referrer-Policy", "no-referrer")
        for name, value in (headers or {}).items():
            self.send_header(name, value)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, value: Any, status: int = 200) -> None:
        self._send(json.dumps(value, ensure_ascii=False).encode("utf-8"),
                   "application/json; charset=utf-8", status)

    def _payload(self) -> Dict[str, Any]:
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError as exc:
            raise WebError("The request size was invalid.") from exc
        if length < 0 or length > MAX_BODY:
            raise WebError("The request is too large.", 413)
        if self.headers.get_content_type() != "application/json":
            raise WebError("The request must use JSON.", 415)
        try:
            value = json.loads(self.rfile.read(length).decode("utf-8")) if length else {}
        except (UnicodeDecodeError, ValueError) as exc:
            raise WebError("The request data was invalid.") from exc
        if not isinstance(value, dict):
            raise WebError("The request data was invalid.")
        return value

    def do_GET(self) -> None:
        if not self._allowed():
            self._json({"error": "This server accepts localhost requests only."}, 403)
            return
        try:
            parsed = urlparse(self.path)
            if parsed.path == "/api/state":
                self._json(self.app.state_json())
                return
            if parsed.path == "/api/runtime":
                self._json({"app": APP_NAME, "pid": os.getpid()})
                return
            if parsed.path == "/api/file":
                query = parse_qs(parsed.query)
                kind = query.get("kind", ["preview"])[0]
                path = self.app.export_pdf if kind == "export" else self.app.preview_pdf
                if not path or not path.is_file():
                    raise WebError("Generate the PDF first.", 404)
                disposition = "attachment" if query.get("download") == ["1"] else "inline"
                self._send(path.read_bytes(), "application/pdf", headers={
                    "Content-Disposition": '%s; filename="%s"' % (disposition, path.name)
                })
                return
            self._static(parsed.path)
        except WebError as exc:
            self._json({"error": str(exc)}, exc.status)
        except (BrokenPipeError, ConnectionResetError, socket.timeout):
            return
        except Exception:
            self._json({"error": "The local server hit an unexpected error."}, 500)

    def _static(self, path_text: str) -> None:
        relative = "index.html" if path_text in {"", "/"} else path_text.lstrip("/")
        path = (WEB_ROOT / relative).resolve()
        if WEB_ROOT not in path.parents and path != WEB_ROOT:
            raise WebError("File not found.", 404)
        if not path.is_file():
            raise WebError("File not found.", 404)
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        if mime.startswith("text/") or mime in {"application/javascript", "application/json"}:
            mime += "; charset=utf-8"
        self._send(path.read_bytes(), mime)

    def do_POST(self) -> None:
        if not self._allowed():
            self._json({"error": "This server accepts localhost requests only."}, 403)
            return
        try:
            payload = self._payload()
            path = urlparse(self.path).path
            if path == "/api/import":
                result = self.app.import_upload(payload)
            elif path == "/api/latest":
                result = self.app.import_latest()
            elif path == "/api/sheet":
                result = self.app.switch_sheet(
                    str(payload.get("sheet", "")),
                    bool(payload.get("excludeHidden", True)),
                )
            elif path == "/api/workspace":
                result = self.app.save_workspace(payload.get("workspace"))
            elif path == "/api/render":
                result = self.app.render(payload.get("workspace", {}), False)
            elif path == "/api/preview":
                result = self.app.preview(
                    payload.get("workspace", {}),
                    int(payload.get("page", 0) or 0),
                )
            elif path == "/api/export":
                result = self.app.render(payload.get("workspace", {}), True)
            elif path == "/api/open":
                result = self.app.open_pdf(str(payload.get("kind", "export")))
            elif path == "/api/email":
                result = self.app.email(payload)
            elif path == "/api/shutdown":
                result = {"stopping": True}
                threading.Thread(target=self.server.shutdown, daemon=True).start()
            else:
                raise WebError("Unknown local endpoint.", 404)
            self._json(result)
        except WebError as exc:
            self._json({"error": str(exc)}, exc.status)
        except (BrokenPipeError, ConnectionResetError, socket.timeout):
            return
        except Exception:
            self._json({"error": "The local server hit an unexpected error."}, 500)


class Server(ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = True

    def __init__(self, address: tuple[str, int], app: Application) -> None:
        super().__init__(address, Handler)
        self.app = app


def main() -> int:
    parser = argparse.ArgumentParser(description="Run the Password Slip Generator web workspace.")
    parser.add_argument("--host", default="127.0.0.1", choices=("127.0.0.1", "localhost"))
    parser.add_argument("--port", type=int, default=8766)
    parser.add_argument("--no-open", action="store_true")
    args = parser.parse_args()
    if not 1024 <= args.port <= 65535:
        raise WebError("Choose a port between 1024 and 65535.")
    app = Application()
    url = "http://127.0.0.1:%d/" % args.port
    try:
        server = Server((args.host, args.port), app)
    except OSError as exc:
        if exc.errno in {48, 98} and not args.no_open:
            webbrowser.open(url)
            return 0
        raise
    print("Password Slip Generator is ready at %s" % url, flush=True)
    print("Keep this window open. Press Control-C to stop.", flush=True)
    if not args.no_open:
        threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever(poll_interval=0.25)
    except KeyboardInterrupt:
        print("\nPassword Slip Generator stopped.")
    finally:
        server.server_close()
    return 0


def cli() -> None:
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130)
    except (OSError, WebError) as exc:
        print("Could not start Password Slip Generator: %s" % exc)
        raise SystemExit(2)


if __name__ == "__main__":
    cli()
