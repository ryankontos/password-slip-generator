#!/usr/bin/env python3
"""Generate cut-aligned password slip PDFs from the newest Excel file in Downloads."""

from __future__ import annotations

import json
import math
import os
import re
import shlex
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional
from urllib.parse import quote

from openpyxl import load_workbook
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


APP_NAME = "password-slip-generator"
AUTHOR = "Ryan Kontos"
YEAR = "2026"
LICENSE_NAME = "0BSD"
ROOT_DIR = Path(__file__).resolve().parent.parent
SETTINGS_DIR = ROOT_DIR / "settings"
LAYOUT_FILE = SETTINGS_DIR / "layout_settings.json"
SETTINGS_FILE = SETTINGS_DIR / "settings.json"
MM = 72 / 25.4
ENV_FILE = ROOT_DIR / ".env"


def load_env_file(path: Path) -> None:
    if not path.is_file():
        return
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return

    for line in lines:
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:].lstrip()
        key, separator, value = line.partition("=")
        key = key.strip()
        if not separator or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key):
            continue
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            if value[0] == '"':
                try:
                    value = json.loads(value)
                except (TypeError, ValueError):
                    value = value[1:-1]
            else:
                value = value[1:-1]
        else:
            value = re.split(r"\s+#", value, maxsplit=1)[0].rstrip()
        os.environ.setdefault(key, value)


load_env_file(ENV_FILE)

SUMMARY_TITLE_HEIGHT_MM = 15.0
SUMMARY_HEADER_HEIGHT_MM = 7.0
SUMMARY_ROW_HEIGHT_MM = 6.5
SUMMARY_CELL_PADDING_MM = 1.2
SUMMARY_HEADER_FONT_PT = 8.5
SUMMARY_DATA_FONT_PT = 8.0


@dataclass
class Settings:
    workbook: str = ""
    sheet: str = ""
    columns: list[str] = field(default_factory=list)
    column_numbers: list[int] = field(default_factory=list)
    password_column_numbers: list[int] = field(default_factory=list)
    truncate_column_numbers: list[int] = field(default_factory=list)
    saved_row_filter_sets: list[dict[str, object]] = field(default_factory=list)
    selected_row_filters: list[dict[str, object]] = field(default_factory=list)
    manual_row_numbers: list[int] = field(default_factory=list)
    blank_slips: int = 0
    email_draft: bool = False
    extra_summary_columns: list[str] = field(default_factory=list)
    email_address: str = ""
    output_folder: str = ""
    input_folder: str = ""
    workbook_extensions: list[str] = field(default_factory=lambda: [".xlsx", ".xlsm"])

    header_height_mm: float = 18.0
    data_height_mm: float = 18.0
    slip_padding_mm: float = 0.0
    top_margin_mm: float = 10.0
    bottom_margin_mm: float = 10.0
    side_margin_mm: float = 0.0
    column_gap_mm: float = 2.0
    padding_mm: float = 2.0
    cut_tick_mm: float = 4.0
    column_width_evenness: float = 0.55
    column_min_width_ratio: float = 0.07
    column_max_width_ratio: float = 0.45

    header_color: str = "#1769AA"
    data_font: str = "Helvetica-Bold"
    password_font: str = "Courier"
    header_font_pt: float = 12.0
    data_font_pt: float = 15.0
    minimum_font_pt: float = 4.0

    show_footer: bool = True
    show_sheet_name: bool = True
    show_generated_datetime: bool = True
    show_page_numbers: bool = True
    footer_font_pt: float = 7.5
    footer_color: str = "#555555"

    @property
    def slip_height_mm(self) -> float:
        return self.header_height_mm + self.data_height_mm + self.slip_padding_mm


def downloads_folder() -> Path:
    return Path.home() / "Downloads"


def newest_workbook(folder: str, extensions: list[str]) -> Optional[Path]:
    folder = Path(folder).expanduser()
    if not folder.is_dir():
        return None

    allowed = {extension.lower() for extension in extensions}
    files = [
        file for file in folder.iterdir()
        if file.is_file()
        and file.suffix.lower() in allowed
        and not file.name.startswith((".", "~$"))
    ]
    return max(files, key=lambda file: file.stat().st_mtime) if files else None


def read_saved_settings() -> Settings:
    settings = Settings()
    ensure_json_files(settings)
    apply_saved_app_settings(settings)
    apply_saved_layout(settings)
    apply_env_settings(settings)
    save_layout_file(settings)
    settings.input_folder = settings.input_folder or str(downloads_folder())
    settings.output_folder = settings.output_folder or str(downloads_folder())
    return settings


def ensure_json_files(defaults: Settings) -> None:
    SETTINGS_DIR.mkdir(exist_ok=True)
    if not LAYOUT_FILE.exists():
        save_layout_file(defaults)
    if not SETTINGS_FILE.exists():
        save_app_settings(defaults)


def apply_saved_app_settings(settings: Settings) -> None:
    try:
        data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    except (OSError, TypeError, ValueError):
        save_app_settings(settings)
        data = {}

    settings.input_folder = str(data.get("input_folder") or "~/Downloads")
    settings.output_folder = str(downloads_folder())
    extensions = data.get("workbook_extensions") or [".xlsx", ".xlsm"]
    settings.workbook_extensions = [str(extension).strip().lower() for extension in extensions]
    settings.column_numbers = saved_column_numbers(data, "column_letters", "column_numbers")
    settings.password_column_numbers = saved_column_numbers(
        data, "password_column_letters", "password_column_numbers"
    )
    settings.truncate_column_numbers = saved_column_numbers(
        data, "truncate_column_letters", "truncate_column_numbers"
    )
    settings.saved_row_filter_sets = saved_row_filter_sets(data)
    settings.selected_row_filters = clean_row_filters(data.get("selected_row_filters", []))
    settings.blank_slips = clean_whole_number(data.get("blank_slips", 0), 0)
    settings.email_draft = clean_bool(
        data.get("email_draft", data.get("email_pdf", False)),
        False,
    )
    settings.email_address = configured_email_address()


def apply_saved_layout(settings: Settings) -> None:
    try:
        data = json.loads(LAYOUT_FILE.read_text(encoding="utf-8"))
        for key, value in data.items():
            if key in layout_field_names():
                setattr(settings, key, clean_layout_value(settings, key, value))
    except (OSError, TypeError, ValueError):
        print(f"Could not read {LAYOUT_FILE.name}; using built-in layout defaults.")


def apply_env_settings(settings: Settings) -> None:
    """Apply optional PASSWORD_SLIPS_* values over JSON settings."""
    settings.input_folder = env_text("INPUT_FOLDER", settings.input_folder)
    settings.output_folder = env_text("OUTPUT_FOLDER", str(downloads_folder()))
    settings.workbook_extensions = env_json_list(
        "WORKBOOK_EXTENSIONS", settings.workbook_extensions
    )
    settings.column_numbers = env_column_numbers(
        "COLUMN_LETTERS", settings.column_numbers, "COLUMN_NUMBERS"
    )
    settings.password_column_numbers = env_column_numbers(
        "PASSWORD_COLUMN_LETTERS",
        settings.password_column_numbers,
        "PASSWORD_COLUMN_NUMBERS",
    )
    settings.truncate_column_numbers = env_column_numbers(
        "TRUNCATE_COLUMN_LETTERS",
        settings.truncate_column_numbers,
        "TRUNCATE_COLUMN_NUMBERS",
    )
    settings.saved_row_filter_sets = env_row_filter_sets(settings.saved_row_filter_sets)
    settings.selected_row_filters = clean_row_filters(
        env_json("SELECTED_ROW_FILTERS", settings.selected_row_filters)
    )
    settings.blank_slips = env_int("BLANK_SLIPS", settings.blank_slips)
    settings.email_address = configured_email_address()
    settings.extra_summary_columns = configured_extra_summary_columns()

    for name in layout_field_names():
        value = os.environ.get(f"PASSWORD_SLIPS_{name.upper()}")
        if value is not None:
            setattr(settings, name, clean_layout_value(settings, name, value))


def env_text(name: str, default: str) -> str:
    value = os.environ.get(f"PASSWORD_SLIPS_{name}")
    return default if value is None else value.strip()


def env_json(name: str, default):
    value = os.environ.get(f"PASSWORD_SLIPS_{name}")
    if value is None:
        return default
    try:
        return json.loads(value)
    except (TypeError, ValueError):
        return default


def env_json_list(name: str, default: list[str]) -> list[str]:
    value = env_json(name, default)
    if not isinstance(value, list):
        return default
    return [str(item).strip().lower() for item in value if str(item).strip()]


def env_column_numbers(name: str, default: list[int], legacy_name: str) -> list[int]:
    env_name = f"PASSWORD_SLIPS_{name}"
    if env_name in os.environ:
        return clean_column_values(env_json(name, default))
    return clean_column_values(env_json(legacy_name, default))


def env_int(name: str, default: int) -> int:
    value = os.environ.get(f"PASSWORD_SLIPS_{name}")
    return clean_whole_number(value, default) if value is not None and value.strip() else default


def clean_layout_value(settings: Settings, key: str, value):
    current = getattr(settings, key)
    if key.endswith("_color"):
        try:
            HexColor(str(value))
            return str(value)
        except (TypeError, ValueError):
            return current
    if isinstance(current, bool):
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.strip().lower() in {"true", "yes", "1"}:
            return True
        if isinstance(value, str) and value.strip().lower() in {"false", "no", "0"}:
            return False
        return current
    if isinstance(current, float):
        try:
            return float(value)
        except (TypeError, ValueError):
            return current
    if isinstance(current, str):
        return str(value)
    return value


def save_settings(settings: Settings) -> None:
    save_app_settings(settings)


def save_app_settings(settings: Settings) -> None:
    data = {
        "_help": app_settings_help(),
        "input_folder": settings.input_folder or "~/Downloads",
        "workbook_extensions": settings.workbook_extensions,
        "column_letters": column_letters_from_numbers(settings.column_numbers),
        "password_column_letters": column_letters_from_numbers(settings.password_column_numbers),
        "truncate_column_letters": column_letters_from_numbers(settings.truncate_column_numbers),
        "saved_row_filter_sets": serialized_row_filter_sets(settings.saved_row_filter_sets),
        "selected_row_filters": serialized_row_filters(settings.selected_row_filters),
        "blank_slips": settings.blank_slips,
        "email_draft": settings.email_draft,
    }
    SETTINGS_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")


def save_layout_file(settings: Settings) -> None:
    layout = {"_help": layout_settings_help()}
    for name in layout_field_names():
        layout[name] = getattr(settings, name)
    LAYOUT_FILE.write_text(json.dumps(layout, indent=2), encoding="utf-8")


def clean_number_list(value) -> list[int]:
    numbers = []
    for item in value if isinstance(value, list) else []:
        try:
            number = int(item)
        except (TypeError, ValueError):
            continue
        if number > 0 and number not in numbers:
            numbers.append(number)
    return numbers


def saved_column_numbers(data: dict, new_key: str, legacy_key: str) -> list[int]:
    value = data[new_key] if new_key in data else data.get(legacy_key, [])
    return clean_column_values(value)


def clean_column_values(value) -> list[int]:
    numbers = []
    for item in value if isinstance(value, list) else []:
        number = column_number_from_text(str(item))
        if number and number not in numbers:
            numbers.append(number)
    return numbers


def column_letters_from_numbers(numbers: list[int]) -> list[str]:
    return [column_letter(number) for number in clean_number_list(numbers)]


def clean_whole_number(value, default: int = 0) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return default
    return max(0, number)


def clean_bool(value, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        value = value.strip().lower()
        if value in {"true", "yes", "1", "y", "on"}:
            return True
        if value in {"false", "no", "0", "n", "off"}:
            return False
    return default


def configured_email_address() -> str:
    for key in ("PASSWORD_SLIPS_EMAIL_ADDRESS", "EMAIL_ADDRESS"):
        value = os.environ.get(key, "").strip()
        if value:
            return value
    return ""


def configured_extra_summary_columns() -> list[str]:
    value = env_json("EXTRA_SUMMARY_COLUMNS", [])
    if not isinstance(value, list):
        return []
    columns = []
    for item in value:
        title = str(item).strip()
        if not title:
            continue
        if title not in columns:
            columns.append(title)
    return columns


def app_settings_help() -> dict[str, str]:
    return {
        "input_folder": "Folder searched for the latest Excel workbook when you press Enter at the file prompt.",
        "workbook_extensions": "Excel file extensions to look for in input_folder.",
        "column_letters": "Last selected column letters, in the order they should appear on each slip.",
        "password_column_letters": "Selected column letters that should use password_font. Add * after a column letter when choosing columns.",
        "truncate_column_letters": "Selected column letters that may truncate instead of shrinking text. Add - after a column letter when choosing columns.",
        "saved_row_filter_sets": "Named quick sets of one or more row rules. Rules in a set are chained with AND.",
        "selected_row_filters": "Rules selected on the last run. The matching quick set is marked next time.",
        "blank_slips": "Extra blank slips to add after the automatically filled last slip page. Press Enter at the prompt to reuse this number.",
        "email_draft": "Last answer to the optional post-export email draft prompt. This is remembered by the app rather than configured in .env.",
    }


def layout_settings_help() -> dict[str, str]:
    return {
        "header_height_mm": "Height of the blue header area on each slip.",
        "data_height_mm": "Height of the white data area on each slip.",
        "slip_padding_mm": "White space below each slip before the next slip starts.",
        "top_margin_mm": "Blank space at the top of each A4 page.",
        "bottom_margin_mm": "Footer text baseline spacing from the bottom of each A4 page, also reserved away from slips.",
        "side_margin_mm": "Left and right page margin.",
        "column_gap_mm": "Space between fields across the slip.",
        "padding_mm": "Inner text padding inside each field area.",
        "cut_tick_mm": "Length of small cut marks at the page edges.",
        "column_width_evenness": "How evenly fields share width. 1 is equal widths, 0 follows each row's text lengths.",
        "column_min_width_ratio": "Smallest share of slip width any field should receive.",
        "column_max_width_ratio": "Largest share of slip width any field should receive.",
        "header_color": "Header colour as a hex value.",
        "data_font": "Font used for normal data values.",
        "password_font": "Font used for columns marked with * when choosing columns.",
        "header_font_pt": "Maximum header text size.",
        "data_font_pt": "Maximum data text size.",
        "minimum_font_pt": "Smallest text size allowed when fitting long text.",
        "show_footer": "Show a small footer in the bottom page margin without changing slip positions.",
        "show_sheet_name": "Show the Excel sheet name in the footer.",
        "show_generated_datetime": "Show the generated date and time in the footer.",
        "show_page_numbers": "Show page numbers as Page X of Y in the footer.",
        "footer_font_pt": "Footer text size.",
        "footer_color": "Footer text colour as a hex value.",
    }


def layout_field_names() -> tuple[str, ...]:
    return (
        "header_height_mm",
        "data_height_mm",
        "slip_padding_mm",
        "top_margin_mm",
        "bottom_margin_mm",
        "side_margin_mm",
        "column_gap_mm",
        "padding_mm",
        "cut_tick_mm",
        "column_width_evenness",
        "column_min_width_ratio",
        "column_max_width_ratio",
        "header_color",
        "data_font",
        "password_font",
        "header_font_pt",
        "data_font_pt",
        "minimum_font_pt",
        "show_footer",
        "show_sheet_name",
        "show_generated_datetime",
        "show_page_numbers",
        "footer_font_pt",
        "footer_color",
    )


def workbook_sheet_names(path: str) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def workbook_headers(path: str, sheet_name: str) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        row = next(workbook[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True), ())
        return unique_labels(row)
    finally:
        workbook.close()


def workbook_records(path: str, sheet_name: str, columns: list[str],
                     row_filters: Optional[list[dict[str, object]]] = None,
                     row_numbers: Optional[list[int]] = None) -> list[list[str]]:
    if not columns:
        return []

    wanted_rows = set(row_numbers or [])
    # Row visibility is stored on worksheet dimensions, which read-only mode
    # does not expose.
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = unique_labels(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        indexes = [headers.index(column) for column in columns]
        records = []
        if wanted_rows:
            row_numbers = sorted(wanted_rows)
        else:
            row_numbers = None

        for row_number, row in worksheet_rows(sheet, row_numbers, max(indexes) + 1):
            if wanted_rows and row_number not in wanted_rows:
                continue
            if not wanted_rows and not row_matches_filters(row, row_filters):
                continue
            values = ["" if index >= len(row) or row[index] is None else str(row[index]) for index in indexes]
            if wanted_rows or any(value.strip() for value in values):
                records.append(values)
        return records
    finally:
        workbook.close()


def worksheet_rows(sheet, row_numbers, max_col: int):
    if row_numbers is None:
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row_is_hidden(sheet, row_number):
                yield row_number, row
        return

    for row_number in row_numbers:
        if row_is_hidden(sheet, row_number):
            continue
        row = next(
            sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=max_col, values_only=True),
            (),
        )
        yield row_number, row


def row_is_hidden(sheet, row_number: int) -> bool:
    dimension = sheet.row_dimensions.get(row_number)
    return bool(dimension and dimension.hidden)


def workbook_column_values(path: str, sheet_name: str, column_number: int) -> list[str]:
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[sheet_name]
        values = []
        seen = set()
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, min_col=column_number, max_col=column_number, values_only=True),
            start=2,
        ):
            if row_is_hidden(sheet, row_number):
                continue
            value = "" if row[0] is None else str(row[0]).strip()
            if value and value not in seen:
                values.append(value)
                seen.add(value)
        return values
    finally:
        workbook.close()


def unique_labels(values) -> list[str]:
    labels = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        label = str(value).strip() if value is not None else f"Column {index}"
        counts[label] = counts.get(label, 0) + 1
        labels.append(label if counts[label] == 1 else f"{label} ({counts[label]})")
    return labels


def clean_row_filters(value) -> list[dict[str, object]]:
    filters = []
    for item in value if isinstance(value, list) else []:
        rule = clean_row_filter(item)
        if rule and rule_key(rule) not in {rule_key(existing) for existing in filters}:
            filters.append(rule)
    return filters


def clean_row_filter_sets(value) -> list[dict[str, object]]:
    sets = []
    for item in value if isinstance(value, list) else []:
        if not isinstance(item, dict):
            continue
        rules = clean_row_filters(item.get("rules", []))
        if not rules:
            continue
        name = str(item.get("name", "")).strip()
        rule_set = {"name": name, "rules": rules}
        if row_filter_set_key(rule_set) not in {
            row_filter_set_key(existing) for existing in sets
        }:
            sets.append(rule_set)
    return sets


def saved_row_filter_sets(data: dict) -> list[dict[str, object]]:
    if "saved_row_filter_sets" in data:
        return clean_row_filter_sets(data["saved_row_filter_sets"])
    legacy_rules = clean_row_filters(data.get("row_filters", []))
    return [
        {"name": f"Saved rule {index}", "rules": [rule]}
        for index, rule in enumerate(legacy_rules, start=1)
    ]


def env_row_filter_sets(default: list[dict[str, object]]) -> list[dict[str, object]]:
    if "PASSWORD_SLIPS_SAVED_ROW_FILTER_SETS" in os.environ:
        return clean_row_filter_sets(env_json("SAVED_ROW_FILTER_SETS", default))
    if "PASSWORD_SLIPS_ROW_FILTERS" in os.environ:
        legacy_rules = clean_row_filters(env_json("ROW_FILTERS", []))
        return [
            {"name": f"Saved rule {index}", "rules": [rule]}
            for index, rule in enumerate(legacy_rules, start=1)
        ]
    return clean_row_filter_sets(default)


def clean_row_filter(value) -> dict[str, object]:
    if not isinstance(value, dict):
        return {}

    mode = str(value.get("mode", "include")).strip().lower()
    if mode not in {"include", "exclude"}:
        mode = "include"

    column_value = value.get("column_letter", value.get("column_number", 0))
    column_number = column_number_from_text(str(column_value))
    if column_number is None:
        return {}

    text = str(value.get("value", "")).strip()
    if column_number < 1 or not text:
        return {}

    return {
        "mode": mode,
        "column_number": column_number,
        "value": text,
    }


def serialized_row_filters(filters: list[dict[str, object]]) -> list[dict[str, object]]:
    serialized = []
    for rule in clean_row_filters(filters):
        output = dict(rule)
        output["column_letter"] = column_letter(int(output.pop("column_number")))
        serialized.append(output)
    return serialized


def serialized_row_filter_sets(value: list[dict[str, object]]) -> list[dict[str, object]]:
    return [
        {
            "name": str(rule_set.get("name", "")).strip(),
            "rules": serialized_row_filters(rule_set.get("rules", [])),
        }
        for rule_set in clean_row_filter_sets(value)
    ]


def row_matches_filters(row, row_filters: Optional[list[dict[str, object]]]) -> bool:
    rules = clean_row_filters(row_filters)
    if not rules:
        return True
    return all(
        row_matches_rule(row, rule)
        if rule["mode"] == "include"
        else not row_matches_rule(row, rule)
        for rule in rules
    )


def row_matches_rule(row, rule: dict[str, object]) -> bool:
    index = int(rule["column_number"]) - 1
    cell = "" if index >= len(row) or row[index] is None else str(row[index])
    return cell.strip() == str(rule["value"])


def rule_key(rule: dict[str, object]) -> tuple[str, int, str]:
    return (
        str(rule.get("mode", "include")),
        int(rule.get("column_number", 0)),
        str(rule.get("value", "")),
    )


def row_filter_set_key(rule_set: dict[str, object]) -> tuple[tuple[str, int, str], ...]:
    return tuple(rule_key(rule) for rule in clean_row_filters(rule_set.get("rules", [])))


def slips_per_page(settings: Settings) -> int:
    usable_height = 297 - settings.top_margin_mm - settings.bottom_margin_mm
    if settings.slip_height_mm <= 0 or usable_height <= 0:
        return 0
    return int(usable_height // settings.slip_height_mm)


def page_count(settings: Settings, slip_count: int) -> int:
    per_page = slips_per_page(settings)
    return math.ceil(slip_count / per_page) if per_page else 0


def column_widths(columns: list[str], record: list[str], settings: Settings, available_width: float) -> list[float]:
    if not columns:
        return []

    if len(columns) == 1:
        return [available_width]

    even_width = available_width / len(columns)
    evenness = clamp(settings.column_width_evenness, 0, 1)
    scores = []
    for index, column in enumerate(columns):
        value = record[index] if index < len(record) else ""
        header_score = stringWidth(column, "Helvetica-Bold", settings.header_font_pt)
        font = data_font_for_column(settings, index)
        value_score = stringWidth(value, font, settings.data_font_pt)
        scores.append(max(1, header_score, value_score))

    total = sum(scores)
    text_widths = [available_width * score / total for score in scores]
    widths = [even_width * evenness + text_width * (1 - evenness) for text_width in text_widths]

    minimum = min(even_width, available_width * clamp(settings.column_min_width_ratio, 0, 1))
    maximum = max(minimum, available_width * clamp(settings.column_max_width_ratio, 0, 1))
    return fit_widths(widths, minimum, maximum, available_width)


def fit_widths(widths: list[float], minimum: float, maximum: float, total_width: float) -> list[float]:
    widths = [min(max(width, minimum), maximum) for width in widths]
    difference = total_width - sum(widths)

    for _ in range(len(widths) * 2):
        if abs(difference) < 0.01:
            break
        if difference > 0:
            adjustable = [index for index, width in enumerate(widths) if width < maximum]
        else:
            adjustable = [index for index, width in enumerate(widths) if width > minimum]
        if not adjustable:
            break
        change = difference / len(adjustable)
        for index in adjustable:
            widths[index] = min(max(widths[index] + change, minimum), maximum)
        difference = total_width - sum(widths)

    if widths:
        widths[-1] += total_width - sum(widths)
    return widths


def clamp(value: float, low: float, high: float) -> float:
    return min(high, max(low, float(value)))


def shrink_to_fit(text: str, font: str, maximum: float, minimum: float, width: float, height: float) -> float:
    size = min(maximum, height / 1.15)
    measured = stringWidth(text, font, size)
    if not text or measured <= width:
        return size
    return max(minimum, min(size, width / max(1, stringWidth(text, font, 1))))


def short_text(text: str, font: str, size: float, width: float) -> str:
    if stringWidth(text, font, size) <= width:
        return text
    suffix = "..."
    if stringWidth(suffix, font, size) > width:
        return ""
    while text and stringWidth(text + suffix, font, size) > width:
        text = text[:-1]
    return text + suffix


def draw_pdf_text(pdf: canvas.Canvas, text: str, font: str, maximum: float, minimum: float,
                  x: float, y: float, width: float, height: float, color, shrink: bool = True) -> None:
    if shrink:
        size = shrink_to_fit(text, font, maximum, minimum, width, height)
    else:
        size = min(maximum, height / 1.15)
    pdf.setFillColor(color)
    pdf.setFont(font, size)
    pdf.drawCentredString(x + width / 2, y + (height - size) / 2 + size * 0.18, short_text(text, font, size, width))


def output_path(settings: Settings) -> Path:
    workbook = Path(settings.workbook)
    name = workbook.stem if workbook.name else "password slips"
    return Path(settings.output_folder or downloads_folder()).expanduser() / f"{name} - password slips.pdf"


def selected_data_records(settings: Settings) -> list[list[str]]:
    return workbook_records(
        settings.workbook,
        settings.sheet,
        settings.columns,
        settings.selected_row_filters,
        settings.manual_row_numbers,
    )


def automatic_blank_slips(row_count: int, per_page: int) -> int:
    """Return the blank slips needed to finish the last non-empty slip page."""
    if row_count <= 0 or per_page <= 0:
        return 0
    return (-row_count) % per_page


def blank_records(settings: Settings, count: Optional[int] = None) -> list[list[str]]:
    if count is None:
        count = clean_whole_number(settings.blank_slips, 0)
    return [[""] * len(settings.columns) for _ in range(clean_whole_number(count, 0))]


def generated_records(settings: Settings, data_records: Optional[list[list[str]]] = None) -> list[list[str]]:
    data_records = selected_data_records(settings) if data_records is None else data_records
    per_page = slips_per_page(settings)
    automatic_blanks = automatic_blank_slips(len(data_records), per_page)
    return (
        list(data_records)
        + blank_records(settings, automatic_blanks)
        + blank_records(settings, settings.blank_slips)
    )


def selected_records(settings: Settings) -> list[list[str]]:
    return generated_records(settings)


def summary_rows_per_page(settings: Settings) -> int:
    reserved_height = (
        SUMMARY_TITLE_HEIGHT_MM
        + SUMMARY_HEADER_HEIGHT_MM
        + settings.bottom_margin_mm
        + 2.0
    )
    usable_height = 297 - settings.top_margin_mm - reserved_height
    return max(1, int(usable_height // SUMMARY_ROW_HEIGHT_MM))


def summary_page_count(settings: Settings, row_count: int) -> int:
    if row_count <= 0:
        return 1
    return math.ceil(row_count / summary_rows_per_page(settings))


def summary_column_widths(columns: list[str], records: list[list[str]], settings: Settings,
                          available_width: float) -> list[float]:
    if not columns:
        return []
    if len(columns) == 1:
        return [available_width]

    scores = []
    for index, column in enumerate(columns):
        score = stringWidth(column, "Helvetica-Bold", SUMMARY_HEADER_FONT_PT)
        for record in records:
            value = record[index] if index < len(record) else ""
            score = max(score, stringWidth(value, "Helvetica", SUMMARY_DATA_FONT_PT))
        scores.append(max(1, score))

    even_width = available_width / len(columns)
    text_total = sum(scores)
    text_widths = [available_width * score / text_total for score in scores]
    widths = [even_width * 0.35 + text_width * 0.65 for text_width in text_widths]
    minimum = min(even_width, available_width * 0.05)
    maximum = max(minimum, available_width * 0.45)
    return fit_widths(widths, minimum, maximum, available_width)


def summary_table(settings: Settings, records: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    extra_columns = list(settings.extra_summary_columns)
    columns = list(settings.columns) + extra_columns
    table_records = [
        list(record) + [""] * len(extra_columns)
        for record in records
    ]
    return columns, table_records


def make_pdf(settings: Settings, output: Optional[Path] = None,
             data_records: Optional[list[list[str]]] = None) -> tuple[int, int]:
    data_records = selected_data_records(settings) if data_records is None else data_records
    records = generated_records(settings, data_records)
    if not records:
        raise ValueError("No slips to generate.")

    per_page = slips_per_page(settings)
    if per_page < 1:
        raise ValueError("The slip height and margins do not fit on A4.")

    page_width, page_height = A4
    top = settings.top_margin_mm * MM
    side = settings.side_margin_mm * MM
    header_height = settings.header_height_mm * MM
    data_height = settings.data_height_mm * MM
    slip_height = settings.slip_height_mm * MM
    gap = settings.column_gap_mm * MM
    padding = settings.padding_mm * MM
    content_width = page_width - side * 2
    column_area = content_width - gap * (len(settings.columns) - 1)

    output = output or output_path(settings)
    output.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(output), pagesize=A4, pageCompression=1)
    pdf.setTitle(f"{settings.sheet or APP_NAME} - password slips")

    slip_pages = page_count(settings, len(records))
    summary_pages = summary_page_count(settings, len(data_records))
    pages = slip_pages + summary_pages
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    for summary_page in range(summary_pages):
        draw_summary_page(
            pdf,
            settings,
            data_records,
            summary_page,
            summary_page + 1,
            pages,
            generated_at,
            page_width,
            page_height,
        )
        pdf.showPage()

    for page in range(slip_pages):
        draw_cut_ticks(pdf, settings, page_width, page_height, per_page)
        draw_footer(pdf, settings, summary_pages + page + 1, pages, generated_at, page_width)
        page_records = records[page * per_page:(page + 1) * per_page]
        for slot, record in enumerate(page_records):
            widths = column_widths(settings.columns, record, settings, column_area)
            draw_slip(pdf, settings, record, widths, side, page_height - top - slot * slip_height, content_width)
        pdf.showPage()

    pdf.save()
    if not output.is_file():
        raise ValueError(f"Could not create PDF at {output}")
    return len(records), pages


def open_email_draft(email_address: str, subject: str) -> bool:
    if not email_address.strip():
        return False
    mailto = f"mailto:{quote(email_address, safe='@')}?subject={quote(subject)}"
    try:
        open_default_application(mailto)
        return True
    except (OSError, subprocess.SubprocessError):
        return False


def open_default_application(target: str) -> None:
    """Open a URL or file with the operating system's default application."""
    if sys.platform.startswith("win"):
        os.startfile(target)  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.run(["open", target], check=True)
    else:
        subprocess.run(["xdg-open", target], check=True)


def email_subject(settings: Settings, generated_at: Optional[datetime] = None) -> str:
    generated_at = generated_at or datetime.now()
    sheet = settings.sheet or "Workbook"
    return f"Generated Password Slips: {sheet} — {generated_at:%Y-%m-%d %H:%M}"


def draw_cut_ticks(pdf: canvas.Canvas, settings: Settings, page_width: float, page_height: float, per_page: int) -> None:
    top = settings.top_margin_mm * MM
    slip_height = settings.slip_height_mm * MM
    tick = settings.cut_tick_mm * MM
    pdf.setStrokeColor(HexColor("#777777"))
    pdf.setLineWidth(0.35)
    for boundary in range(per_page + 1):
        y = page_height - top - boundary * slip_height
        pdf.line(0, y, tick, y)
        pdf.line(page_width - tick, y, page_width, y)


def draw_footer(pdf: canvas.Canvas, settings: Settings, page_number: int, page_total: int,
                generated_at: str, page_width: float) -> None:
    if not settings.show_footer:
        return

    parts = []
    if settings.show_sheet_name and settings.sheet:
        parts.append(f"Sheet: {settings.sheet}")
    if settings.show_generated_datetime:
        parts.append(f"Generated: {generated_at}")
    if settings.show_page_numbers:
        parts.append(f"Page {page_number} of {page_total}")
    if not parts:
        return

    side = settings.side_margin_mm * MM
    y = max(3 * MM, settings.bottom_margin_mm * MM)
    text = "  |  ".join(parts)
    available_width = page_width - side * 2
    size = shrink_to_fit(text, "Helvetica", settings.footer_font_pt, settings.minimum_font_pt,
                         available_width, settings.footer_font_pt * 1.5)
    pdf.setFillColor(HexColor(settings.footer_color))
    pdf.setFont("Helvetica", size)
    pdf.drawCentredString(page_width / 2, y, short_text(text, "Helvetica", size, available_width))


def draw_summary_page(pdf: canvas.Canvas, settings: Settings, records: list[list[str]],
                      summary_page: int, page_number: int, page_total: int,
                      generated_at: str, page_width: float, page_height: float) -> None:
    side = settings.side_margin_mm * MM
    content_width = page_width - side * 2
    top = page_height - settings.top_margin_mm * MM
    title = f"{settings.sheet or 'Workbook'} - Summary"

    pdf.setFillColor(HexColor("#111111"))
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(side, top - 7 * MM, short_text(title, "Helvetica-Bold", 16, content_width))
    pdf.setFillColor(HexColor("#555555"))
    pdf.setFont("Helvetica", 9)
    pdf.drawString(side, top - 12 * MM, f"Generated: {generated_at}")

    table_top = top - SUMMARY_TITLE_HEIGHT_MM * MM
    header_height = SUMMARY_HEADER_HEIGHT_MM * MM
    row_height = SUMMARY_ROW_HEIGHT_MM * MM
    padding = SUMMARY_CELL_PADDING_MM * MM
    table_columns, table_records = summary_table(settings, records)
    widths = summary_column_widths(table_columns, table_records, settings, content_width)
    rows_per_page = summary_rows_per_page(settings)
    page_records = table_records[summary_page * rows_per_page:(summary_page + 1) * rows_per_page]

    if not table_columns:
        pdf.setFillColor(HexColor("#333333"))
        pdf.setFont("Helvetica", 10)
        pdf.drawString(side, table_top - 10 * MM, "No columns were selected.")
        draw_footer(pdf, settings, page_number, page_total, generated_at, page_width)
        return

    header_y = table_top - header_height
    x = side
    pdf.setStrokeColor(HexColor("#B8B8B8"))
    pdf.setLineWidth(0.45)
    for index, width in enumerate(widths):
        pdf.setFillColor(HexColor("#E9EDF1"))
        pdf.rect(x, header_y, width, header_height, stroke=1, fill=1)
        draw_summary_cell(
            pdf,
            table_columns[index],
            "Helvetica-Bold",
            SUMMARY_HEADER_FONT_PT,
            x + padding,
            header_y,
            max(1, width - padding * 2),
            header_height,
        )
        x += width

    if not page_records:
        pdf.setFillColor(HexColor("#333333"))
        pdf.setFont("Helvetica", 10)
        pdf.drawString(side + padding, header_y - 10 * MM, "No workbook rows were selected.")
    else:
        for row_index, record in enumerate(page_records):
            row_y = header_y - (row_index + 1) * row_height
            x = side
            for column_index, width in enumerate(widths):
                value = record[column_index] if column_index < len(record) else ""
                pdf.setFillColor(HexColor("#FFFFFF" if row_index % 2 == 0 else "#F7F7F7"))
                pdf.rect(x, row_y, width, row_height, stroke=1, fill=1)
                draw_summary_cell(
                    pdf,
                    value,
                    "Helvetica",
                    SUMMARY_DATA_FONT_PT,
                    x + padding,
                    row_y,
                    max(1, width - padding * 2),
                    row_height,
                )
                x += width

    draw_footer(pdf, settings, page_number, page_total, generated_at, page_width)


def draw_summary_cell(pdf: canvas.Canvas, text: str, font: str, maximum: float,
                      x: float, y: float, width: float, height: float) -> None:
    text = str(text)
    size = shrink_to_fit(text, font, maximum, 4.5, width, height)
    pdf.setFillColor(HexColor("#111111"))
    pdf.setFont(font, size)
    pdf.drawString(x, y + (height - size) / 2 + size * 0.18,
                   short_text(text, font, size, width))


def draw_slip(pdf: canvas.Canvas, settings: Settings, record: list[str], widths: list[float],
              left: float, top: float, content_width: float) -> None:
    header_height = settings.header_height_mm * MM
    data_height = settings.data_height_mm * MM
    gap = settings.column_gap_mm * MM
    padding = settings.padding_mm * MM
    header_y = top - header_height
    data_y = header_y - data_height

    pdf.setFillColor(HexColor(settings.header_color))
    pdf.rect(left, header_y, content_width, header_height, stroke=0, fill=1)

    x = left
    for index, width in enumerate(widths):
        text_width = max(1, width - padding * 2)
        draw_pdf_text(pdf, settings.columns[index], "Helvetica-Bold", settings.header_font_pt,
                      settings.minimum_font_pt, x + padding, header_y, text_width, header_height, HexColor("#FFFFFF"))
        draw_pdf_text(pdf, record[index], data_font_for_column(settings, index), settings.data_font_pt, settings.minimum_font_pt,
                      x + padding, data_y, text_width, data_height, HexColor("#000000"),
                      shrink=not column_may_truncate(settings, index))
        x += width + gap


def data_font_for_column(settings: Settings, index: int) -> str:
    column_number = settings.column_numbers[index] if index < len(settings.column_numbers) else index + 1
    font = settings.password_font if column_number in settings.password_column_numbers else settings.data_font
    return usable_font(font, "Helvetica-Bold")


def column_may_truncate(settings: Settings, index: int) -> bool:
    column_number = settings.column_numbers[index] if index < len(settings.column_numbers) else index + 1
    return column_number in settings.truncate_column_numbers


def usable_font(font: str, fallback: str) -> str:
    try:
        stringWidth("test", font, 10)
        return font
    except Exception:
        return fallback


def prompt(label: str, default: str = "") -> str:
    shown = f" [{default}]" if default else ""
    value = input(f"{label}{shown}: ").strip()
    return value or default


def prompt_yes_no(label: str, default: bool = False) -> bool:
    shown_default = "y" if default else "n"
    while True:
        value = prompt(f"{label} (y/n)", shown_default).strip().lower()
        if value in {"y", "yes", "1", "true"}:
            return True
        if value in {"n", "no", "0", "false"}:
            return False
        print("Enter y or n.")


def banner() -> None:
    print()
    print(APP_NAME)
    print(f"Created by {AUTHOR}, {YEAR} | {LICENSE_NAME} licensed")
    print("Cut-aligned A4 password slip PDFs from Excel.")
    print("Press Enter to accept a suggestion.")


def section(title: str) -> None:
    print()
    print(title)
    print("-" * len(title))


def clean_path(value: str) -> str:
    try:
        parts = shlex.split(value, posix=not sys.platform.startswith("win"))
        return parts[0].strip("'\"") if len(parts) == 1 else value
    except ValueError:
        return value.strip("'\"")


def choose_workbook(settings: Settings) -> str:
    section("Workbook")
    latest = newest_workbook(settings.input_folder, settings.workbook_extensions)
    if latest:
        print(f"Latest Excel file in {settings.input_folder}: {latest}")
    else:
        print(f"No Excel files found in {settings.input_folder}.")

    while True:
        default = str(latest) if latest else ""
        label = "Excel file path (Enter = latest file)" if latest else "Excel file path"
        path = Path(clean_path(prompt(label, default))).expanduser()
        if path.is_file() and path.suffix.lower() in set(settings.workbook_extensions):
            settings.input_folder = str(path.parent)
            return str(path)
        allowed = ", ".join(settings.workbook_extensions)
        print(f"That Excel file is not available. Enter a path ending in {allowed}.")


def choose_from_list(title: str, choices: list[str], default: str = "", reverse: bool = False) -> str:
    if not choices:
        raise ValueError(f"No {title.lower()} found.")

    shown_choices = list(reversed(choices)) if reverse else choices
    section(title)
    for index, choice in enumerate(shown_choices, start=1):
        marker = " *" if choice == default else ""
        print(f"  {index}. {choice}{marker}")

    default_index = str(shown_choices.index(default) + 1) if default in shown_choices else "1"
    while True:
        value = prompt(f"Choose {title.lower()} number", default_index)
        if value.isdigit() and 1 <= int(value) <= len(shown_choices):
            return shown_choices[int(value) - 1]
        if value in choices:
            return value
        print("Enter a number from the list.")


def choose_columns(headers: list[str], settings: Settings) -> tuple[list[str], list[int], list[int], list[int]]:
    if not headers:
        raise ValueError("No column headings were found in the first row.")

    default_numbers = remembered_column_numbers(headers, settings)
    default_password_numbers = remembered_password_column_numbers(default_numbers, settings)
    default_truncate_numbers = remembered_truncate_column_numbers(default_numbers, settings)

    section("Columns")
    for index, header in enumerate(headers, start=1):
        marker = " *" if index in default_numbers else ""
        annotations = []
        if index in default_password_numbers:
            annotations.append("password font")
        if index in default_truncate_numbers:
            annotations.append("allows truncation")
        annotation = f" [{', '.join(annotations)}]" if annotations else ""
        print(f"  {column_letter(index)}. {header}{marker}{annotation}")
    print("Enter letters in the order to print them, for example A,C,D.")
    print("Add * for password_font and - to allow truncation, for example A,B*-,C-.")

    while True:
        default = column_letters_for_prompt(default_numbers, default_password_numbers, default_truncate_numbers)
        value = prompt("Column letters", default)
        if value.strip().lower() == "all":
            numbers = list(range(1, len(headers) + 1))
            password_numbers = []
            truncate_numbers = []
        else:
            numbers, password_numbers, truncate_numbers = column_numbers_from_text(value, headers)

        if not numbers:
            print("Enter at least one valid column letter.")
            continue

        selected_headers = [headers[number - 1] for number in numbers]
        print("Selected columns, in print order:")
        for number, header in zip(numbers, selected_headers):
            annotations = []
            if number in password_numbers:
                annotations.append("password font")
            if number in truncate_numbers:
                annotations.append("allows truncation")
            annotation = f" [{', '.join(annotations)}]" if annotations else ""
            print(f"  {column_letter(number)}. {header}{annotation}")

        confirmation = prompt("Use these columns? (y/n)", "y").strip().lower()
        if confirmation in {"", "y", "yes"}:
            return selected_headers, numbers, password_numbers, truncate_numbers
        if confirmation not in {"n", "no"}:
            print("Enter y to confirm or n to choose again.")


def remembered_column_numbers(headers: list[str], settings: Settings) -> list[int]:
    numbers = [
        number for number in settings.column_numbers
        if isinstance(number, int) and 1 <= number <= len(headers)
    ]
    if numbers:
        return numbers

    return list(range(1, len(headers) + 1))


def remembered_password_column_numbers(column_numbers: list[int], settings: Settings) -> list[int]:
    return [
        number for number in settings.password_column_numbers
        if number in column_numbers
    ]


def remembered_truncate_column_numbers(column_numbers: list[int], settings: Settings) -> list[int]:
    return [
        number for number in settings.truncate_column_numbers
        if number in column_numbers
    ]


def column_letters_for_prompt(column_numbers: list[int], password_column_numbers: list[int],
                              truncate_column_numbers: list[int]) -> str:
    return ",".join(
        column_letter(number)
        + ("*" if number in password_column_numbers else "")
        + ("-" if number in truncate_column_numbers else "")
        for number in column_numbers
    )


def column_numbers_from_text(value: str, headers: list[str]) -> tuple[list[int], list[int], list[int]]:
    tokens = value.replace(",", " ").split()
    numbers = []
    password_numbers = []
    truncate_numbers = []
    for token in tokens:
        token, is_password, may_truncate = split_column_token(token)
        number = column_number_from_text(token)
        if number is None:
            return [], [], []
        if not 1 <= number <= len(headers):
            return [], [], []
        if number not in numbers:
            numbers.append(number)
        if is_password and number not in password_numbers:
            password_numbers.append(number)
        if may_truncate and number not in truncate_numbers:
            truncate_numbers.append(number)
    return numbers, password_numbers, truncate_numbers


def split_column_token(token: str) -> tuple[str, bool, bool]:
    token = token.strip()
    is_password = False
    may_truncate = False
    while token and token[-1] in {"*", "-"}:
        marker = token[-1]
        token = token[:-1]
        if marker == "*":
            is_password = True
        if marker == "-":
            may_truncate = True
    return token, is_password, may_truncate


def column_letter(number: int) -> str:
    letters = ""
    while number > 0:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def column_number_from_text(value: str) -> Optional[int]:
    value = value.strip().upper()
    if value.isdigit():
        return int(value)
    if not value.isalpha():
        return None

    number = 0
    for letter in value:
        number = number * 26 + ord(letter) - 64
    return number


def choose_row_filters(headers: list[str], settings: Settings) -> list[dict[str, object]]:
    while True:
        saved = usable_row_filter_sets(headers, settings.saved_row_filter_sets)
        selected_number = selected_row_filter_set_number(saved, settings.selected_row_filters)

        section("Rows")
        print("Choose a saved rule set, create one, or use specific spreadsheet rows.")
        print("  0. All workbook rows")
        for index, rule_set in enumerate(saved, start=1):
            marker = " *" if index == selected_number else ""
            print(f"  {index}. {row_filter_set_description(rule_set, headers)}{marker}")
        print("  n. New rule set")
        print("  r. Specific spreadsheet row numbers")

        default = str(selected_number or 0)
        value = prompt("Row choice", default).strip().lower()
        if value == "0":
            settings.selected_row_filters = []
            settings.manual_row_numbers = []
            return []
        if value in {"new", "n"}:
            rule_set = create_row_filter_set(headers, settings)
            remember_row_filter_set(settings, rule_set)
            settings.selected_row_filters = list(rule_set["rules"])
            settings.manual_row_numbers = []
            return settings.selected_row_filters
        if value in {"rows", "r", "custom", "c"}:
            settings.selected_row_filters = []
            settings.manual_row_numbers = choose_manual_row_numbers()
            return []

        if value.isdigit() and 1 <= int(value) <= len(saved):
            rule_set = saved[int(value) - 1]
            settings.selected_row_filters = list(rule_set["rules"])
            settings.manual_row_numbers = []
            return settings.selected_row_filters

        print("Enter 0, a saved set number, n, or r.")


def choose_blank_slips(settings: Settings, row_count: int) -> int:
    per_page = slips_per_page(settings)
    automatic_blanks = automatic_blank_slips(row_count, per_page)
    section("Blank slips")
    if row_count:
        print(
            f"There will be {automatic_blanks} blank "
            f"{plural(automatic_blanks, 'slip')} on the last page."
        )
        if automatic_blanks:
            print("These blanks will automatically finish the last password-slip page.")
        else:
            print("The selected rows already fill the last password-slip page.")
    else:
        print("No workbook rows matched the current selection.")
        print("Extra blank slips can still be generated if you need them.")
    print("How many extra slips do you want? Extra slips will add page(s).")

    while True:
        value = prompt("Extra blank slips", str(settings.blank_slips)).strip()
        if value == "":
            return clean_whole_number(settings.blank_slips, 0)
        try:
            number = int(value)
        except ValueError:
            print("Enter 0 or a whole number.")
            continue
        if number >= 0:
            return number
        print("Enter 0 or a whole number.")


def usable_row_filter_sets(headers: list[str], value: list[dict[str, object]]) -> list[dict[str, object]]:
    return [
        rule_set for rule_set in clean_row_filter_sets(value)
        if all(int(rule["column_number"]) <= len(headers) for rule in rule_set["rules"])
    ]


def selected_row_filter_set_number(saved: list[dict[str, object]],
                                   selected: list[dict[str, object]]) -> int:
    selected_key = tuple(rule_key(rule) for rule in clean_row_filters(selected))
    if not selected_key:
        return 0
    for index, rule_set in enumerate(saved, start=1):
        if row_filter_set_key(rule_set) == selected_key:
            return index
    return 0


def choose_manual_row_numbers() -> list[int]:
    while True:
        value = prompt("Spreadsheet row numbers to include, separated by commas or ranges").strip()
        numbers = row_numbers_from_text(value)
        if numbers:
            return numbers
        print("Enter row numbers 2 or higher, for example 2,5,9 or 10-15.")


def row_numbers_from_text(value: str) -> list[int]:
    numbers = []
    value = re.sub(r"\s*-\s*", "-", value.strip())
    for token in re.split(r"[\s,]+", value):
        if not token:
            continue
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            if not start_text.isdigit() or not end_text.isdigit():
                return []
            start = int(start_text)
            end = int(end_text)
            if start < 2 or end < start:
                return []
            for number in range(start, end + 1):
                if number not in numbers:
                    numbers.append(number)
            continue
        if not token.isdigit():
            return []
        number = int(token)
        if number < 2:
            return []
        if number not in numbers:
            numbers.append(number)
    return numbers


def create_row_filter_set(headers: list[str], settings: Settings) -> dict[str, object]:
    section("New rule set")
    print("Rules are applied together: every rule must match.")
    rules = []
    while True:
        rule = create_row_rule(headers, settings, len(rules) + 1)
        rules.append(rule)
        print(f"  Added: {row_filter_description(rule, headers)}")
        if not prompt_yes_no("Add another rule", False):
            break

    default_name = " + ".join(short_rule_description(rule, headers) for rule in rules)
    if len(default_name) > 60:
        default_name = default_name[:57] + "..."
    name = prompt("Save this quick set as", default_name or "My rules").strip()
    return {"name": name, "rules": rules}


def create_row_rule(headers: list[str], settings: Settings, rule_number: int) -> dict[str, object]:
    print()
    print(f"Rule {rule_number}")
    for index, header in enumerate(headers, start=1):
        print(f"  {column_letter(index)}. {header}")
    while True:
        column = prompt("Column letter").strip()
        number = column_number_from_text(column)
        if number is not None and 1 <= number <= len(headers):
            break
        print(f"Enter a column from A to {column_letter(len(headers))}.")

    print("  1. is")
    print("  2. is not")
    mode_text = prompt("Comparison", "1").strip().lower()
    mode = "exclude" if mode_text in {"2", "not", "is not", "exclude"} else "include"
    text = choose_filter_value(settings, number)
    return {
        "mode": mode,
        "column_number": number,
        "value": text,
    }


def choose_filter_value(settings: Settings, column_number: int) -> str:
    values = workbook_column_values(settings.workbook, settings.sheet, column_number)

    if values:
        shown_values = values[:20]
        print("Choose a value number, or type a value directly:")
        for index, value in enumerate(shown_values, start=1):
            shown = value if len(value) <= 70 else value[:67] + "..."
            print(f"  {index}. {shown}")
        if len(values) > len(shown_values):
            print(f"  ...and {len(values) - len(shown_values)} more; type one directly if needed.")

        choice = prompt("Value", "1").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(shown_values):
            return shown_values[int(choice) - 1]
        if choice:
            return choice
    else:
        print("No existing values found in that column.")

    while True:
        text = prompt("Custom text to match exactly").strip()
        if text:
            return text
        print("Enter the text to match.")


def remember_row_filter_set(settings: Settings, rule_set: dict[str, object]) -> None:
    cleaned = clean_row_filter_sets([rule_set])
    if not cleaned:
        return
    new_set = cleaned[0]
    saved = [
        existing for existing in clean_row_filter_sets(settings.saved_row_filter_sets)
        if row_filter_set_key(existing) != row_filter_set_key(new_set)
    ]
    saved.insert(0, new_set)
    settings.saved_row_filter_sets = saved[:20]


def row_filter_description(rule: dict[str, object], headers: list[str]) -> str:
    rule = clean_row_filter(rule)
    number = int(rule["column_number"])
    header = headers[number - 1] if 1 <= number <= len(headers) else f"Column {column_letter(number)}"
    action = "is" if rule["mode"] == "include" else "is not"
    value = str(rule["value"])
    return f'{column_letter(number)} ({header}) {action} "{value}"'


def short_rule_description(rule: dict[str, object], headers: list[str]) -> str:
    rule = clean_row_filter(rule)
    number = int(rule["column_number"])
    header = headers[number - 1] if 1 <= number <= len(headers) else column_letter(number)
    comparison = "=" if rule["mode"] == "include" else "≠"
    return f'{header} {comparison} {rule["value"]}'


def row_filter_set_description(rule_set: dict[str, object], headers: list[str]) -> str:
    cleaned = clean_row_filter_sets([rule_set])
    if not cleaned:
        return "Invalid rule set"
    rule_set = cleaned[0]
    name = str(rule_set.get("name", "")).strip()
    rules = " AND ".join(
        row_filter_description(rule, headers) for rule in rule_set["rules"]
    )
    return f"{name}: {rules}" if name else rules


def choose_email_draft(settings: Settings) -> bool:
    section("Email")
    print(f"Create a draft email to {settings.email_address}?")
    print("The PDF has been exported and can be attached from the path shown above.")
    return prompt_yes_no("Open draft", settings.email_draft)


def preview_records(columns: list[str], records: list[list[str]]) -> None:
    section("Preview")
    if not records:
        print("  No data rows found with those columns.")
        return

    for row_number, row in enumerate(records[:3], start=1):
        print(f"  Slip {row_number}:")
        for column, value in zip(columns, row):
            shown = value if len(value) <= 70 else value[:67] + "..."
            print(f"    {column}: {shown}")
    if len(records) > 3:
        print(f"  ...and {len(records) - 3} more")


def preview_blank_slips(automatic_count: int, extra_count: int) -> None:
    if automatic_count:
        print(
            f"  Plus {automatic_count} automatic blank "
            f"{plural(automatic_count, 'slip')} to finish the last page."
        )
    if extra_count:
        print(f"  Plus {extra_count} extra blank {plural(extra_count, 'slip')} on new page(s).")


def print_summary(settings: Settings, row_count: int, slip_count: int) -> None:
    per_page = slips_per_page(settings)
    automatic_blanks = automatic_blank_slips(row_count, per_page)
    slip_pages = page_count(settings, slip_count)
    summary_pages = summary_page_count(settings, row_count)
    print()
    print(f"Ready: {row_count} workbook {plural(row_count, 'row')}.")
    print(
        f"  Password slips: {slip_count} slips across "
        f"{slip_pages} {plural(slip_pages, 'page')} ({per_page} per page)."
    )
    if automatic_blanks:
        print(f"  Automatic blanks: {automatic_blanks} to finish the last slip page.")
    if settings.blank_slips:
        print(f"  Extra blanks: {settings.blank_slips} on new page(s).")
    if summary_pages:
        print(f"  Summary: {summary_pages} compact {plural(summary_pages, 'page')}.")


def plural(count: int, singular: str) -> str:
    return singular if count == 1 else singular + "s"


def run_cli() -> None:
    banner()

    settings = read_saved_settings()

    while True:
        settings.workbook = choose_workbook(settings)
        try:
            sheets = workbook_sheet_names(settings.workbook)
            break
        except Exception as error:
            print(f"Could not open that workbook: {error}")
            settings.workbook = ""

    settings.sheet = choose_from_list("Sheets", sheets, settings.sheet, reverse=True)

    headers = workbook_headers(settings.workbook, settings.sheet)
    (
        settings.columns,
        settings.column_numbers,
        settings.password_column_numbers,
        settings.truncate_column_numbers,
    ) = choose_columns(headers, settings)
    settings.selected_row_filters = choose_row_filters(headers, settings)

    workbook_rows = workbook_records(
        settings.workbook,
        settings.sheet,
        settings.columns,
        settings.selected_row_filters,
        settings.manual_row_numbers,
    )
    per_page = slips_per_page(settings)
    if per_page < 1:
        raise ValueError("The slip height and margins do not fit on A4.")

    settings.blank_slips = choose_blank_slips(settings, len(workbook_rows))
    automatic_blanks = automatic_blank_slips(len(workbook_rows), per_page)
    records = generated_records(settings, workbook_rows)
    preview_records(settings.columns, workbook_rows)
    preview_blank_slips(automatic_blanks, settings.blank_slips)
    print_summary(settings, len(workbook_rows), len(records))
    if not records:
        raise ValueError("No slips to generate. Choose matching rows or add blank slips.")

    if not Path(settings.workbook).is_file():
        raise ValueError("Choose an Excel workbook.")
    if not settings.columns:
        raise ValueError("Choose at least one column.")

    save_settings(settings)
    count, pages = make_pdf(settings, data_records=workbook_rows)
    pdf = output_path(settings)

    print()
    print("Done")
    print(f"  Created {count} slips across {pages} {plural(pages, 'page')}.")
    print("  The first page(s) contain the compact summary.")
    print(f"  PDF: {pdf}")

    if settings.email_address:
        settings.email_draft = choose_email_draft(settings)
        save_settings(settings)
        if settings.email_draft:
            subject = email_subject(settings)
            if open_email_draft(settings.email_address, subject):
                print(f"  Opened a draft addressed to {settings.email_address}.")
                print(f"  Subject: {subject}")
            else:
                print("  Could not open the default mail app.")
        else:
            print("  Email draft skipped.")


def main() -> None:
    try:
        run_cli()
    except KeyboardInterrupt:
        print("\nCancelled.")
        sys.exit(130)
    except Exception as error:
        print(f"\nError: {error}")
        sys.exit(1)


if __name__ == "__main__":
    main()
